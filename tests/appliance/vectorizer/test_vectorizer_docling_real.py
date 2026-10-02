"""Real Docling: these tests run only where Docling is installed.

The repository's virtualenv does not have it, so there they are skipped. They
are run with an interpreter that has the Docling of the image (the versions in
appliance/vectorizer/requirements.lock), pytest and cryptography:

    <venv with docling>/bin/python -m pytest tests/appliance/vectorizer -q -p no:cacheprovider

Documents are generated here with the libraries Docling itself depends on
(python-docx, python-pptx, openpyxl). The office and web formats need no
model. The PDF pipeline needs layout and table models (downloaded from
Hugging Face, or found in DOCLING_ARTIFACTS_PATH as in the image); where they
are not available the PDF test checks that this is reported as
`parser_unavailable` and is then skipped, not passed.
"""

from __future__ import annotations

import importlib.metadata
import logging
import re
from pathlib import Path

import pytest

pytest.importorskip("docling")

from hm_vectorizer.chunker import CHUNK_MAX_CHARS
from hm_vectorizer.parsers import DefaultParser, ParseError
from hm_vectorizer.selfcheck import minimal_pdf
from hm_vectorizer.sync import run_sync
from vz_fakes import FakeOllama, FakeQdrant
from vz_support import VECTORIZER_DIR, Site

COLLECTION = "happymining_docs"


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Maintenance handbook", level=1)
    document.add_paragraph("This handbook describes how the cooling loop of container C2 is serviced.")
    document.add_heading("Pump replacement", level=2)
    for step in range(1, 7):
        document.add_paragraph(
            f"Step {step}: close valve V{step}, wait until the gauge reads zero, then remove the four bolts. "
            "The replacement pump reference is HM-PUMP-7731. " * 6
        )
    document.add_heading("Contacts", level=2)
    table = document.add_table(rows=3, cols=2)
    for row, (role, phone) in enumerate(
        [("Role", "Extension"), ("Site manager", "4410"), ("Electrician", "4417")]
    ):
        table.cell(row, 0).text = role
        table.cell(row, 1).text = phone
    document.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    slide.placeholders[1].text = "Hashrate grew by twelve percent in the third quarter."
    presentation.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Container", "Power kW"])
    sheet.append(["C1", 410])
    sheet.append(["C2", 395])
    workbook.save(path)


def make_pdf(path: Path, lines: list[str]) -> None:
    """A one-page PDF with a text layer (the one the image's self-check converts)."""
    path.write_bytes(minimal_pdf(lines))


@pytest.fixture(scope="module")
def parser() -> DefaultParser:
    return DefaultParser(ocr=False, max_file_bytes=64 * 1024 * 1024)


def test_installed_docling_is_the_version_the_image_installs() -> None:
    """Otherwise these tests say nothing about the image."""
    lock = (VECTORIZER_DIR / "requirements.lock").read_text(encoding="utf-8")
    for name in ("docling", "docling-core", "docling-parse", "docling-ibm-models"):
        pinned = re.search(rf"^{name}==(\S+) ", lock, flags=re.M)
        assert pinned is not None, name
        assert importlib.metadata.version(name) == pinned.group(1), name


def test_formats_let_through(parser: DefaultParser) -> None:
    for ext in ("pdf", "docx", "pptx", "xlsx", "html", "htm", "csv", "adoc", "md", "txt"):
        assert parser.unsupported_reason(ext) is None, ext
    for ext in ("mp3", "mp4", "wav", "zip", "exe", "png", "jpg"):
        assert parser.unsupported_reason(ext) == "unsupported_format", ext


def test_images_are_let_through_only_with_ocr() -> None:
    with_ocr = DefaultParser(ocr=True, max_file_bytes=1024)
    assert with_ocr.unsupported_reason("png") is None and with_ocr.unsupported_reason("jpg") is None
    assert with_ocr.unsupported_reason("mp3") == "unsupported_format"


def test_docx_is_parsed_and_cut_with_headings(tmp_path: Path, parser: DefaultParser) -> None:
    path = tmp_path / "doc.docx"
    make_docx(path)
    chunks = parser.parse(path, "docx")
    assert len(chunks) >= 6
    assert all(0 < len(c.text) <= CHUNK_MAX_CHARS for c in chunks)
    assert chunks[0].headings == ("Maintenance handbook",)
    assert "cooling loop of container C2" in chunks[0].text
    steps = [c for c in chunks if c.headings == ("Maintenance handbook", "Pump replacement")]
    assert len(steps) >= 4 and any("Step 3: close valve V3" in c.text for c in steps)
    contacts = [c for c in chunks if c.headings == ("Maintenance handbook", "Contacts")]
    assert contacts and "4410" in contacts[0].text and "Electrician" in contacts[0].text
    assert chunks[1].embed_text.startswith("Maintenance handbook\nPump replacement\n")


def test_pptx_xlsx_html_csv_are_parsed(tmp_path: Path, parser: DefaultParser) -> None:
    make_pptx(tmp_path / "doc.pptx")
    make_xlsx(tmp_path / "doc.xlsx")
    (tmp_path / "doc.html").write_text(
        "<html><body><h1>Safety rules</h1><p>Wear ear protection inside the containers.</p></body></html>"
    )
    (tmp_path / "doc.csv").write_text("miner,hashrate\nS21,200\nS19,110\n")
    text = {
        ext: "\n".join(c.embed_text for c in parser.parse(tmp_path / f"doc.{ext}", ext))
        for ext in ("pptx", "xlsx", "html", "csv")
    }
    assert "twelve percent" in text["pptx"]
    assert "395" in text["xlsx"] and "C2" in text["xlsx"]
    assert "Wear ear protection" in text["html"] and "Safety rules" in text["html"]
    assert "S21" in text["csv"] and "200" in text["csv"]


def test_damaged_document_fails_with_a_code(tmp_path: Path, parser: DefaultParser) -> None:
    path = tmp_path / "doc.docx"
    path.write_bytes(b"PK\x03\x04 this is not a real archive" * 50)
    with pytest.raises(ParseError) as caught:
        parser.parse(path, "docx")
    assert caught.value.code == "parse_error"


def test_html_does_not_fetch_what_it_references(tmp_path: Path, parser: DefaultParser) -> None:
    """Docling's backends do not fetch remote or local resources by default; this checks
    that a document pointing at a local server causes no request to it."""
    hit = FakeOllama()
    try:
        (tmp_path / "doc.html").write_text(
            f'<html><body><p>Text with an image.</p><img src="{hit.url}/pixel.png">'
            f'<link rel="stylesheet" href="{hit.url}/style.css"><iframe src="{hit.url}/frame"></iframe>'
            "</body></html>"
        )
        chunks = parser.parse(tmp_path / "doc.html", "html")
        assert any("Text with an image." in c.text for c in chunks)
        assert hit.requests == []
    finally:
        hit.close()


def test_pdf_with_the_layout_pipeline(tmp_path: Path, parser: DefaultParser) -> None:
    path = tmp_path / "doc.pdf"
    make_pdf(path, ["Invoice 2026-0042", "Customer: Example SARL", "Total due: 1250.00 EUR"])
    try:
        chunks = parser.parse(path, "pdf")
    except ParseError as exc:
        # Models missing (not downloaded, no network) is reported as such, never as a broken file.
        assert exc.code == "parser_unavailable"
        pytest.skip("Docling's PDF pipeline could not be prepared here: its models are not available")
    text = "\n".join(c.text for c in chunks)
    assert "Invoice 2026-0042" in text and "1250.00" in text
    assert all(c.page == 1 for c in chunks)


def test_sync_with_real_documents(
    site: Site, ollama: FakeOllama, qdrant: FakeQdrant, caplog: pytest.LogCaptureFixture
) -> None:
    docs = site.source_dir()
    (docs / "zebra-okapi").mkdir()
    make_docx(docs / "zebra-okapi" / "narwhal-handbook.docx")
    make_pptx(docs / "zebra-okapi" / "platypus-review.pptx")
    make_xlsx(docs / "power.xlsx")
    (docs / "broken-wombat.docx").write_bytes(b"PK\x03\x04 not an archive" * 50)
    (docs / "notes.md").write_text("# Notes\n\nPlain Markdown still goes through the direct reader.")
    for path in docs.rglob("*"):
        if path.is_file():
            site.write(str(path.relative_to(docs)), path.read_bytes())  # sets a modification time in the past

    with caplog.at_level(logging.INFO):
        result = run_sync(site.load(), site.state_dir)
    assert result.status["state"] == "idle"
    assert (result.status["files_indexed"], result.status["files_failed"]) == (4, 1)
    assert result.status["detail"] == "failed: parse_error=1"
    assert qdrant.paths(COLLECTION) == {
        "zebra-okapi/narwhal-handbook.docx",
        "zebra-okapi/platypus-review.pptx",
        "power.xlsx",
        "notes.md",
    }
    handbook = [p for p in qdrant.payloads(COLLECTION) if p["path"].endswith("handbook.docx")]
    assert any(p["headings"] == ["Maintenance handbook", "Pump replacement"] for p in handbook)

    logged = "\n".join(r.getMessage() for r in caplog.records if r.levelno >= logging.INFO)
    for forbidden in ("zebra", "okapi", "narwhal", "platypus", "wombat", "HM-PUMP", "twelve percent"):
        assert forbidden not in logged, "neither our logs nor Docling's name a file or quote a document"

    ollama.clear()
    again = run_sync(site.load(), site.state_dir)
    assert (again.processed, again.unchanged) == (0, 4)
    assert ollama.embedded_texts() == []
